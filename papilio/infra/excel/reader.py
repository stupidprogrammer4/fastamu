import asyncio
from concurrent.futures import ProcessPoolExecutor
from multiprocessing import get_context
from typing import Any

from openpyxl import load_workbook

from .row import ExcelRow


class ExcelReader:
    """Async Excel reader — the counterpart to `ExcelWriter`.

    Reads run in a separate process via a ``ProcessPoolExecutor`` (a queue
    fronting worker process(es)) and are ``await``-ed. ``read_rows`` maps each
    sheet row onto a typed `ExcelRow` by field order, so pydantic validates and
    coerces the cell types in the worker. Row classes must be importable at
    module scope, and their values must be pickleable. ``close()`` on app
    shutdown.
    """

    def __init__(self, max_workers: int = 1) -> None:
        self._pool = ProcessPoolExecutor(
            max_workers=max_workers, mp_context=get_context("spawn")
        )

    @staticmethod
    def _read_rows_job[TRow: ExcelRow](
        path: str,
        sheet: str | None,
        start_row: int,
        row_model: type[TRow],
        max_rows: int | None,
    ) -> list[TRow]:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            if ws is None:
                raise ValueError("workbook has no active worksheet")

            names = row_model.column_names
            out: list[TRow] = []
            for i, row in enumerate(
                ws.iter_rows(
                    min_row=start_row, max_col=len(names), values_only=True
                )
            ):
                if max_rows is not None and i >= max_rows:
                    break
                if all(v is None for v in row):  # stop at the first blank row
                    break
                out.append(row_model(**dict(zip(names, row))))
            return out
        finally:
            wb.close()

    @staticmethod
    def _read_cell_job(path: str, sheet: str | None, cell: str) -> Any:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet] if sheet else wb.active
            if ws is None:
                raise ValueError("workbook has no active worksheet")
            return ws[cell].value
        finally:
            wb.close()

    async def read_rows[TRow: ExcelRow](
        self,
        path: str,
        row_model: type[TRow],
        *,
        start_row: int,
        sheet: str | None = None,
        limit: int | None = None,
    ) -> list[TRow]:
        """Read rows from ``start_row`` (1-based) into ``row_model`` instances,
        mapping columns to fields by order. Stops at the first blank row or
        after ``limit`` rows. ``row_model`` must be importable by the worker;
        returned models must be pickleable."""
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(
            self._pool,
            self._read_rows_job,
            path,
            sheet,
            start_row,
            row_model,
            limit,
        )

    async def read_cell(
        self, path: str, cell: str, *, sheet: str | None = None
    ) -> Any:
        """Read a cell value using a read-only workbook in the worker.

        Rows are scanned up to the requested cell; later rows cost more.
        Formula cells return their cached value, without recalculation.
        """
        loop = asyncio.get_running_loop()
        result = await loop.run_in_executor(
            self._pool, self._read_cell_job, path, sheet, cell
        )
        return result

    async def close(self) -> None:
        await asyncio.to_thread(self._pool.shutdown)
