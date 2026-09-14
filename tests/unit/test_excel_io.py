import os
from datetime import datetime
from multiprocessing import current_process
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from pydantic import PrivateAttr, ValidationError

from papilio.infra.excel.reader import ExcelReader
from papilio.infra.excel.row import ExcelRow, Row
from papilio.infra.excel.writer import ExcelWriter


class ProductRow(ExcelRow):
    name: str = Row(title="Name")
    count: int = Row(title="Count")


class WorkerRow(ProductRow):
    _pid: int = PrivateAttr(default_factory=os.getpid)

    @classmethod
    def titles(cls) -> list[str]:
        assert current_process().name != "MainProcess"
        return super().titles()

    def cells(self) -> list:
        assert os.getpid() != self._pid
        return super().cells()


async def test_read_cell_preserves_values_and_sheet_selection(tmp_path):
    path = tmp_path / "cells.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "first sheet"
    sheet = workbook.create_sheet("values")
    sheet["A1"] = "active sheet"
    sheet["B2"] = datetime(2026, 1, 2)
    sheet["C3"] = "merged"
    sheet.merge_cells("C3:D3")
    sheet["E4"] = "=1+2"
    sheet["F4"] = "=2+3"
    workbook.active = 1
    workbook.save(path)
    workbook.close()
    # openpyxl does not calculate formulas; simulate a saved Excel result.
    with ZipFile(path) as source:
        entries = [(info, source.read(info)) for info in source.infolist()]
    with ZipFile(path, "w") as target:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet2.xml":
                root = ElementTree.fromstring(data)
                value = root.find(
                    ".//s:c[@r='E4']/s:v",
                    {
                        "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                    },
                )
                assert value is not None
                value.text = "3"
                data = ElementTree.tostring(root)
            target.writestr(info, data)

    reader = ExcelReader()
    try:
        assert await reader.read_cell(str(path), "A1") == "active sheet"
        assert (
            await reader.read_cell(str(path), "A1", sheet="Sheet")
            == "first sheet"
        )
        assert await reader.read_cell(str(path), "B2") == datetime(2026, 1, 2)
        assert await reader.read_cell(str(path), "C3") == "merged"
        assert await reader.read_cell(str(path), "D3") is None
        assert await reader.read_cell(str(path), "Z100") is None
        assert await reader.read_cell(str(path), "E4") == 3
        assert await reader.read_cell(str(path), "F4") is None
        with pytest.raises(KeyError):
            await reader.read_cell(str(path), "A1", sheet="missing")
        assert await reader.read_cell(str(path), "A1") == "active sheet"
    finally:
        await reader.close()


async def test_excel_roundtrip(tmp_path):
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook.save(template)
    workbook.close()
    writer, reader = ExcelWriter(), ExcelReader()
    try:
        await writer.write_rows(
            str(template),
            str(output),
            [ProductRow(name="پروانه", count=2)],
            start_row=1,
            with_titles=True,
        )
        rows = await reader.read_rows(str(output), ProductRow, start_row=2)
        assert rows == [ProductRow(name="پروانه", count=2)]
        assert await reader.read_cell(str(output), "A1") == "Name"
        await writer.write_cell(str(output), str(output), "B2", 3)
        assert await reader.read_cell(str(output), "B2") == 3
    finally:
        await writer.close()
        await reader.close()


async def test_excel_rows_are_converted_in_worker_processes(tmp_path):
    template, output = tmp_path / "template.xlsx", tmp_path / "output.xlsx"
    workbook = Workbook()
    workbook.save(template)
    workbook.close()
    writer, reader = ExcelWriter(), ExcelReader()
    try:
        await writer.write_rows(
            str(template),
            str(output),
            [WorkerRow(name="item", count=2)],
            start_row=1,
            with_titles=True,
        )
        rows = await reader.read_rows(str(output), WorkerRow, start_row=2)
        assert len(rows) == 1
        assert rows[0].model_dump() == {"name": "item", "count": 2}
        assert rows[0]._pid != os.getpid()
        assert await reader.read_cell(str(output), "A1") == "Name"
    finally:
        await writer.close()
        await reader.close()


async def test_excel_worker_preserves_selection_and_validation(tmp_path):
    path = tmp_path / "rows.xlsx"
    workbook = Workbook()
    sheet = workbook.create_sheet("items")
    sheet.append(["Name", "Count"])
    sheet.append(["first", "2"])
    sheet.append(["second", 3])
    sheet.append([None, None])
    sheet.append(["invalid", "not a number"])
    workbook.save(path)
    workbook.close()
    reader = ExcelReader()
    try:
        rows = await reader.read_rows(
            str(path),
            ProductRow,
            sheet="items",
            start_row=2,
            limit=1,
        )
        assert rows == [ProductRow(name="first", count=2)]
        rows = await reader.read_rows(
            str(path),
            ProductRow,
            sheet="items",
            start_row=2,
        )
        assert rows == [
            ProductRow(name="first", count=2),
            ProductRow(name="second", count=3),
        ]
        with pytest.raises(ValidationError) as error:
            await reader.read_rows(
                str(path),
                ProductRow,
                sheet="items",
                start_row=5,
            )
        assert error.value.errors()[0]["loc"] == ("count",)
        assert error.value.errors()[0]["type"] == "int_parsing"
        # A failed job must leave the pool usable.
        assert (
            await reader.read_rows(
                str(path),
                ProductRow,
                sheet="items",
                start_row=2,
                limit=0,
            )
            == []
        )
    finally:
        await reader.close()
